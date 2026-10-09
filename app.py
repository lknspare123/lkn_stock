"""
app.py — LKN Stock web app (Flask + Neon Postgres + Cloudinary), deployable on Render.
Web port of the desktop (Tkinter) application: same pages, same Khmer labels, same business rules.
"""
import io
import os
import re
import json
import hmac
import time
from datetime import datetime, timedelta

from flask import (Flask, render_template, request, redirect, url_for, session, flash, jsonify,
                   abort, send_file, send_from_directory)

from core.neon_db import DB
from core import storage
from core import backup as bk
from utils.helpers import (fmt_currency, fmt_qty, fmt_date, safe_float, today_str, calc_profit, now_local,
                           MONEY_EPS, GENERAL_CUSTOMER, payment_state, payment_label, outstanding_balance,
                           compute_discounted_total, fmt_pct)

APP_VERSION = "1.0.5"
APP_AUTHOR = "NITH VANNOY"
ITEMS_SESSION_SECONDS = 30 * 60        # the Items page re-asks for the password after 30 minutes idle (as on desktop)

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-only-change-me")
app.config["MAX_CONTENT_LENGTH"] = 64 * 1024 * 1024     # backups can be large
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=14)
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["JSON_SORT_KEYS"] = False
app.json.ensure_ascii = False
app.json.sort_keys = False

application = app  # lets Render's default "gunicorn app:application" work too

db = DB()

NAV = [
    ("ផ្ទាំងគ្រប់គ្រង", "dashboard", "dashboard", "dashboard"),
    ("ទិន្នន័យមុខទំនិញ", "items", "items", "item"),
    ("ស្តុកចម្បង", "stock", "stock", "stock"),
    ("ទទួលស្តុក", "receive_list", "receive_list", "receive"),
    ("បោះពុម្ពស្លាក", "labels", "labels", "label"),
    ("ការលក់", "sales_list", "sales_list", "sale"),
    ("ផ្ទេរទំនិញសាខា", "transfers_list", "transfers_list", "transfer"),
    ("របាយការណ៍", "reports", "reports", "report"),
    ("ការកំណត់", "settings", "settings", "setting"),
]

app.jinja_env.filters["currency"] = fmt_currency
app.jinja_env.filters["qty"] = fmt_qty
app.jinja_env.filters["dmy"] = fmt_date
app.jinja_env.filters["thumb"] = storage.thumb
app.jinja_env.filters["pay_label"] = lambda row_total, paid: payment_label(row_total, paid)


@app.url_defaults
def static_cache_bust(endpoint, values):
    """Add ?v=<file time> to every /static/ link, so browsers always load the newest JS/CSS right after a deploy
    (no more 'I updated the site but still see the old page')."""
    if endpoint == "static" and "filename" in values:
        try:
            values["v"] = int(os.path.getmtime(os.path.join(app.static_folder, values["filename"])))
        except OSError:
            pass


@app.context_processor
def inject_nav():
    return {"nav": NAV, "current": request.endpoint or "", "today": today_str(), "GENERAL_CUSTOMER": GENERAL_CUSTOMER,
            "app_version": APP_VERSION}


# ── Auth ─────────────────────────────────────────────────────────────────────
def check_password(given):
    expected = os.environ.get("APP_PASSWORD", "")
    return bool(expected) and hmac.compare_digest((given or "").encode(), expected.encode())


def items_password_ok(given):
    expected = os.environ.get("ITEMS_PASSWORD") or os.environ.get("APP_PASSWORD", "")
    return bool(expected) and hmac.compare_digest((given or "").encode(), expected.encode())


@app.before_request
def require_login():
    if request.endpoint in ("login", "static", "healthz", "favicon", "webmanifest", "cron_backup"):
        return
    if not session.get("auth"):
        if request.path.startswith(("/api/", "/items/api/")):
            return jsonify(error="Session expired — please sign in again."), 401
        return redirect(url_for("login", next=request.path))
    # The Items page (and its editing API) is behind a second 30-minute unlock, like the desktop AuthGate.
    if request.path.startswith("/items") and request.endpoint not in ("items_unlock",):
        last = session.get("items_auth_at", 0)
        if time.time() - last > ITEMS_SESSION_SECONDS:
            if request.path.startswith("/items/api/"):
                return jsonify(error="ការអនុញ្ញាតផុតកំណត់ សូមបញ្ចូលពាក្យសម្ងាត់ម្តងទៀត។", locked=True), 401
            return redirect(url_for("items_unlock", next=request.full_path.rstrip("?")))
        session["items_auth_at"] = time.time()


@app.get("/favicon.ico")
def favicon():
    """Browsers ask for /favicon.ico by default, before they read the <link> tags."""
    return send_from_directory(app.static_folder, "favicon.ico", mimetype="image/x-icon", max_age=86400)


@app.get("/site.webmanifest")
def webmanifest():
    return send_file(os.path.join(app.root_path, "site.webmanifest"), mimetype="application/manifest+json", max_age=3600)


@app.get("/healthz")
def healthz():
    return "ok"


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if check_password(request.form.get("password", "")):
            session.clear()
            session["auth"] = True   # the Items page has its own, separate unlock (asked on first visit)
            session.permanent = True
            nxt = request.args.get("next", "")
            return redirect(nxt if nxt.startswith("/") and not nxt.startswith("//") else url_for("dashboard"))
        flash("ពាក្យសម្ងាត់មិនត្រឹមត្រូវ (Wrong password).", "error")
    return render_template("login.html")


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/items/unlock", methods=["GET", "POST"])
def items_unlock():
    nxt = request.values.get("next", "")
    nxt = nxt if nxt.startswith("/items") and not nxt.startswith("//") else url_for("items")
    if request.method == "POST":
        if items_password_ok(request.form.get("password", "")):
            session["items_auth_at"] = time.time()
            return redirect(nxt)
        flash("ពាក្យសម្ងាត់មិនត្រឹមត្រូវ។", "error")
    return render_template("items_unlock.html", next=nxt)


# ── Helpers ──────────────────────────────────────────────────────────────────
def f(name, default=""):
    return (request.form.get(name) or default).strip()


def body():
    return request.get_json(force=True, silent=True) or {}


def err(msg, code=400):
    return jsonify(ok=False, error=msg), code


def items_payload(items):
    """Everything the client-side item search needs (receive / sales / transfers / labels)."""
    out = []
    for i in items:
        names = []
        for x in i["distributors"]:
            n = (x["distributor_name"] or "").strip()
            if n and n not in names:
                names.append(n)
        dates = [x["last_price_date"] for x in i["distributors"] if x.get("last_price_date")]
        mn, mx = i["min_unit_price"], i["max_unit_price"]
        out.append({
            "id": i["id"], "code": i["item_code"], "kh": i.get("name_khmer", ""), "en": i.get("name_english", ""),
            "sell_price": i["sell_price"], "unit_price": i["unit_price"], "min_cost": i["min_unit_price"] or 0,
            "remain": i["remain_qty"], "image": storage.thumb(i.get("image_url", ""), 120),
            "floor": i.get("floor", ""), "corner": i.get("corner", ""),
            "suppliers": ", ".join(names), "cost_lo": mn, "cost_hi": mx,
            "price_date": max(dates) if dates else "",
            "distributors": [{"name": x["distributor_name"], "cost": x["cost_price"]} for x in i["distributors"]],
            "customers": [{"name": x["customer_name"], "price": x["special_sell_price"]} for x in i["customer_prices"]],
        })
    return out


def parse_date(s, default=None):
    s = (s or "").strip()
    try:
        return datetime.strptime(s, "%Y-%m-%d").strftime("%Y-%m-%d")
    except ValueError:
        return default


def month_start():
    return today_str()[:8] + "01"


def pay_row(total, paid):
    total, paid = safe_float(total), safe_float(paid)
    return {"total": total, "paid": paid, "balance": outstanding_balance(total, paid),
            "state": payment_state(total, paid), "label": payment_label(total, paid)}


def combine_datetime(date_str, keep_time_from=None):
    """Picked date + current time of day (or the original time when editing), like the desktop app."""
    d = parse_date(date_str, today_str())
    t = now_local().strftime("%H:%M:%S")
    if keep_time_from and len(keep_time_from) >= 19:
        t = keep_time_from[11:19]
    return f"{d} {t}"


def tel_display(p1, p2):
    p1, p2 = (p1 or "").strip() or "–", (p2 or "").strip()
    return f"{p1} / {p2}" if p2 else p1


def split_tel(display):
    if not display or display == "–":
        return "", ""
    parts = display.split(" / ", 1)
    p1 = "" if parts[0].strip() == "–" else parts[0].strip()
    return p1, (parts[1].strip() if len(parts) == 2 else "")


# ── Dashboard ────────────────────────────────────────────────────────────────
def build_chart(daily):
    """Server-side SVG bar chart geometry (no JS / no external libraries)."""
    W, H, PB, PT = 720, 180, 22, 10
    n = len(daily)
    mx = max([v for _, v in daily] + [0]) or 1
    bw = W / n
    bars, ticks = [], []
    for i, (d, v) in enumerate(daily):
        h = (H - PB - PT) * v / mx
        bars.append({"x": round(i * bw + bw * 0.18, 1), "w": round(bw * 0.64, 1),
                     "y": round(H - PB - max(h, 1.5 if v else 0), 1), "h": round(max(h, 1.5 if v else 0), 1),
                     "tip": f"{fmt_date(d)} — {fmt_currency(v)}"})
        if i % 5 == 0 or i == n - 1:
            ticks.append({"x": round(i * bw + bw / 2, 1), "label": d[5:].replace("-", "/")})
    return {"w": W, "h": H, "base": H - PB, "bars": bars, "ticks": ticks, "max": fmt_currency(mx)}


@app.get("/")
def dashboard():
    d = db.dashboard_data(force=bool(request.args.get("refresh")))
    best_max = max([b["qty"] for b in d["best"]] + [1])
    return render_template("dashboard.html", d=d, chart=build_chart(d["daily"]), best_max=best_max)


# ── Items ────────────────────────────────────────────────────────────────────
def item_row(i):
    sup = []
    for x in i["distributors"]:
        n = (x["distributor_name"] or "").strip()
        if n and n not in sup:
            sup.append(n)
    return {
        "id": i["id"], "code": i["item_code"], "kh": i.get("name_khmer", ""), "en": i.get("name_english", ""),
        "floor": i.get("floor", ""), "corner": i.get("corner", ""),
        "u0": i["min_unit_price"], "u1": i["max_unit_price"], "s0": i["min_sell_price"], "s1": i["max_sell_price"],
        "suppliers": ", ".join(sup), "original": i["original_qty"], "added": i["stock_add_qty"],
        "on_hand": i["on_hand_qty"], "sold": i["sell_qty"], "transferred": i["transfer_qty"], "remain": i["remain_qty"],
        "image": i.get("image_url", ""), "sell_price": i["sell_price"], "unit_price": i["unit_price"],
    }


@app.get("/items")
def items():
    rows = [item_row(i) for i in db.get_all_items()]
    return render_template("items.html", data={"rows": rows, "open_new": bool(request.args.get("new")),
                                               "open_edit": request.args.get("edit", "")})


@app.get("/items/new")
def item_new():
    return redirect(url_for("items", new=1))


@app.get("/items/<item_id>")
def item_edit(item_id):
    return redirect(url_for("items", edit=item_id))


@app.get("/items/api/list")
def items_api_list():
    return jsonify([item_row(i) for i in db.get_all_items()])


@app.post("/items/api/save")
def items_api_save():
    item_id = f("id")
    item = db.get_item(item_id) if item_id else None
    if item_id and not item:
        return err("រកមិនឃើញទំនិញនេះទៀតហើយ។", 404)
    code = item["item_code"] if item else f("item_code")
    kh, en = f("name_khmer"), f("name_english")
    if not code or not kh or not en:
        return err("កូដទំនិញ ឈ្មោះខ្មែរ និងឈ្មោះអង់គ្លេសដាច់ខាតត្រូវតែបំពេញ។")
    url = item.get("image_url", "") if item else ""
    pid = item.get("image_public_id", "") if item else ""
    new_pid = ""
    try:
        file = request.files.get("image")
        if item and request.form.get("remove_image"):
            storage.delete_image(pid)
            url = pid = ""
        if file and file.filename:
            new_url, new_pid = storage.upload_item_image(file, code)
            if pid and pid != new_pid:
                storage.delete_image(pid)
            url, pid = new_url, new_pid
        args = (kh, en, f("floor"), f("corner"), url, pid, safe_float(request.form.get("sell_price")),
                safe_float(request.form.get("unit_price")))
        if item:
            db.update_item(item["id"], *args, original_qty=safe_float(request.form.get("original_qty")))
            new_id = item["id"]
        else:
            new_id = db.insert_item(code, *args, original_qty=safe_float(request.form.get("original_qty")))
    except ValueError as e:
        if new_pid and not item:
            storage.delete_image(new_pid)
        return err(str(e))
    except Exception as e:   # Cloudinary / DB problems
        return err(f"មិនអាចរក្សាទុកទិន្នន័យបានឡើយ៖ {e}", 500)
    return jsonify(ok=True, id=new_id, item=item_row(db.get_item(new_id)))


@app.post("/items/api/<item_id>/delete")
def items_api_delete(item_id):
    item = db.get_item(item_id) or abort(404)
    try:
        db.delete_item(item_id)
    except ValueError as e:
        return err(str(e))
    storage.delete_image(item.get("image_public_id", ""))
    return jsonify(ok=True)


@app.get("/items/api/<item_id>/matrix")
def items_api_matrix(item_id):
    it = db.get_item(item_id) or abort(404)
    return jsonify(
        distributors=[{"name": x["distributor_name"], "price": x["cost_price"], "date": x.get("last_price_date", "")}
                      for x in it["distributors"]],
        customers=[{"name": x["customer_name"], "price": x["special_sell_price"], "date": x.get("last_price_date", "")}
                   for x in it["customer_prices"]],
        suppliers=db.get_all_suppliers(), customer_names=db.get_all_customers())


@app.post("/items/api/<item_id>/<kind>")
def items_api_matrix_save(item_id, kind):
    d = body()
    name = (d.get("name") or "").strip()
    if kind not in ("distributor", "customer"):
        abort(404)
    if not name:
        return err("សូមបំពេញឈ្មោះអ្នកផ្គត់ផ្គង់។" if kind == "distributor" else "សូមបំពេញឈ្មោះអតិថិជន។")
    date = (d.get("date") or "").strip()
    if date and not parse_date(date):
        return err("សូមបញ្ចូលកាលបរិច្ឆេទត្រឹមត្រូវជាទម្រង់ YYYY-MM-DD ឬទុកឲ្យទទេ។")
    date = parse_date(date, today_str())
    price = safe_float(d.get("price"))
    if kind == "distributor":
        db.save_item_distributor(item_id, name, price, date)
    else:
        db.save_item_customer_price(item_id, name, price, date)
    return jsonify(ok=True)


@app.post("/items/api/<item_id>/<kind>/delete")
def items_api_matrix_delete(item_id, kind):
    name = (body().get("name") or "").strip()
    if kind == "distributor":
        db.delete_item_distributor(item_id, name)
    elif kind == "customer":
        db.delete_item_customer_price(item_id, name)
    else:
        abort(404)
    return jsonify(ok=True)


@app.get("/api/items/<item_id>/history")
def item_history(item_id):
    return jsonify([{"date": h.get("receive_date", ""), "invoice": h.get("invoice_code", ""),
                     "distributor": h.get("distributor_name", ""), "qty": h.get("qty_receive", 0),
                     "price": h.get("unit_price", 0)} for h in db.get_item_purchase_history(item_id)])


# ── Main stock (master-detail) ───────────────────────────────────────────────
@app.get("/stock")
def stock():
    rows = []
    for i in db.get_all_items():
        u0, u1 = i["min_unit_price"], i["max_unit_price"]
        s0, s1 = i["min_sell_price"], i["max_sell_price"]
        rem = i["remain_qty"]
        rows.append({
            "id": i["id"], "code": i["item_code"], "kh": i.get("name_khmer", ""), "en": i.get("name_english", ""),
            "floor": i.get("floor", ""), "corner": i.get("corner", ""),
            "original": i["original_qty"], "added": i["stock_add_qty"], "on_hand": i["on_hand_qty"],
            "sold": i["sell_qty"], "transferred": i["transfer_qty"], "remain": rem,
            "u0": u0, "u1": u1, "s0": s0, "s1": s1,
            "total_cost": rem * ((u0 + u1) / 2) if u0 is not None else 0.0,
            "total_sell": rem * ((s0 + s1) / 2) if s0 is not None else 0.0,
            "image": i.get("image_url", ""),
            "distributors": sorted([{"name": x["distributor_name"], "cost": x["cost_price"],
                                     "date": x.get("last_price_date", "")} for x in i["distributors"]],
                                   key=lambda x: x["name"].lower()),
            "customers": sorted([{"name": x["customer_name"], "price": x["special_sell_price"],
                                  "date": x.get("last_price_date", "")} for x in i["customer_prices"]],
                                key=lambda x: x["name"].lower()),
        })
    return render_template("stock.html", data={"rows": rows, "sel": request.args.get("sel", "")})


# ── Stock receive (purchase invoices) ────────────────────────────────────────
def receive_view(r):
    p = pay_row(r["total_cost"], r["paid_amount"])
    return {"code": r["invoice_code"], "item_count": r["item_count"], "date": r["receive_date"],
            "distributor": r["distributor_name"], "notes": r["notes"] or "", **p}


@app.get("/receive")
def receive_list():
    items = db.get_all_items()
    return render_template("receive.html", items=items_payload(items), suppliers=db.get_all_suppliers(),
                           next_code=db.next_code("INV-RCV-", "purchase_invoices", "invoice_code"),
                           view=request.args.get("view", ""))


@app.get("/receive/new")
def receive_new():
    return redirect(url_for("receive_list"))


@app.get("/receive/<code>")
def receive_detail(code):
    return redirect(url_for("receive_list", view=code))


@app.get("/api/receive/next-code")
def api_receive_next_code():
    return jsonify(code=db.next_code("INV-RCV-", "purchase_invoices", "invoice_code"))


@app.get("/api/receive/list")
def api_receive_list():
    return jsonify([receive_view(r) for r in db.get_receive_history()])


@app.get("/api/receive/<code>")
def api_receive_get(code):
    inv = db.get_purchase_invoice(code) or abort(404)
    lines = db.get_invoice_items(code)
    total = sum(l["qty_receive"] * l["unit_price"] for l in lines)
    return jsonify(
        code=inv["invoice_code"], distributor=inv["distributor_name"], date=inv["receive_date"],
        notes=inv["notes"] or "", **pay_row(total, inv["paid_amount"]),
        lines=[{"item_id": l["item_id"], "code": l["item_code"], "name": l["name_khmer"] or l["name_english"] or l["item_code"],
                "qty": l["qty_receive"], "price": l["unit_price"], "total": l["qty_receive"] * l["unit_price"]}
               for l in lines])


@app.post("/api/receive/save")
def api_receive_save():
    d = body()
    code, dist = (d.get("code") or "").strip(), (d.get("distributor") or "").strip()
    if not code or not dist:
        return err("សូមបំពេញលេខវិក្កយបត្រ និងឈ្មោះអ្នកផ្គត់ផ្គង់។")
    lines = [{"item_id": x["item_id"], "qty": safe_float(x["qty"]), "price": safe_float(x["price"])}
             for x in d.get("items", []) if safe_float(x.get("qty")) > 0]
    if not lines:
        return err("សូមបញ្ចូលទំនិញយ៉ាងហោចណាស់មួយមុខទៅក្នុងវិក្កយបត្រ។")
    paid = safe_float(d.get("paid_amount"))
    if paid < 0:
        return err("ប្រាក់ទូទាត់មិនអាចតូចជាងសូន្យបានទេ។")
    editing = (d.get("editing_code") or "").strip() or None
    try:
        db.save_purchase_invoice(code, dist, parse_date(d.get("date"), today_str()), (d.get("notes") or "").strip(),
                                 lines, paid_amount=paid, editing_code=editing)
    except ValueError as e:
        return err(str(e))
    total = sum(l["qty"] * l["price"] for l in lines)
    return jsonify(ok=True, code=code, editing=bool(editing), balance=outstanding_balance(total, paid))


@app.post("/api/receive/<code>/payment")
def api_receive_payment(code):
    d = body()
    r = db.get_purchase_invoice(code) or abort(404)
    new_total = max(safe_float(r["paid_amount"]) + safe_float(d.get("amount")), 0.0)
    db.set_purchase_payment(code, new_total)
    return jsonify(ok=True, **pay_row(r["total_cost_calc"], new_total))


@app.post("/api/receive/mark-paid")
def api_receive_mark_paid():
    return jsonify(ok=True, count=db.mark_purchase_paid(body().get("codes", [])))


@app.post("/api/receive/delete")
def api_receive_delete():
    db.delete_purchase_invoices(body().get("codes", []))
    return jsonify(ok=True)


def summary_rows(rows, total_key, paid_key, code_key, name_key, date_key):
    """Period summary used by the Receive / Sales / Transfer history 'Summary' dialogs (same maths as desktop)."""
    total_all = total_paid = total_unpaid = total_discount = 0.0
    c_unpaid = c_partial = 0
    out = []
    for r in rows:
        amt, paid = safe_float(r[total_key]), safe_float(r[paid_key])
        total_all += amt
        total_paid += min(paid, amt)
        total_unpaid += outstanding_balance(amt, paid)
        total_discount += safe_float(r.get("discount_amount"))
        st = payment_state(amt, paid)
        if st != "paid":
            if st == "unpaid":
                c_unpaid += 1
            else:
                c_partial += 1
            out.append({"code": r[code_key], "name": r[name_key], "date": r[date_key], "total": amt, "paid": paid,
                        "balance": outstanding_balance(amt, paid)})
    return {"count_all": len(rows), "count_partial": c_partial, "count_unpaid": c_unpaid, "total_all": total_all,
            "total_paid": total_paid, "total_unpaid": total_unpaid, "total_discount": total_discount, "rows": out}


def summary_range():
    s, e = parse_date(request.args.get("start"), month_start()), parse_date(request.args.get("end"), today_str())
    return (e, s) if s > e else (s, e)


@app.get("/api/summary/<kind>")
def api_summary(kind):
    s, e = summary_range()
    if kind == "receive":
        rows = [r for r in db.get_receive_history() if s <= (r["receive_date"] or "")[:10] <= e]
        rows = [{**r, "total_cost": r["total_cost"]} for r in rows]
        res = summary_rows(rows, "total_cost", "paid_amount", "invoice_code", "distributor_name", "receive_date")
    elif kind == "sales":
        rows = db.get_sales_history(start=s, end=e)
        res = summary_rows(rows, "grand_total", "paid_amount", "sale_code", "customer_name", "sale_date")
    elif kind == "transfers":
        rows = db.get_transfers_history(start=s, end=e)
        res = summary_rows(rows, "total_cost", "paid_amount", "transfer_code", "branch_name", "created_at")
    else:
        abort(404)
    return jsonify(start=s, end=e, **res)


@app.get("/receive/<code>/print")
def receive_print(code):
    inv = db.get_purchase_invoice(code) or abort(404)
    lines = db.get_invoice_items(code)
    total = sum(l["qty_receive"] * l["unit_price"] for l in lines)
    return render_template("print_receive.html", inv=inv, lines=lines, total=total, pay=pay_row(total, inv["paid_amount"]))


@app.get("/print/summary/<kind>")
def print_summary(kind):
    if kind not in ("receive", "sales", "transfers"):
        abort(404)
    data = api_summary(kind).get_json()
    titles = {"receive": ("សង្ខេបរបាយការណ៍ទទួលស្តុក", "អ្នកផ្គត់ផ្គង់"), "sales": ("សង្ខេបរបាយការណ៍លក់", "អតិថិជន"),
              "transfers": ("សង្ខេបរបាយការណ៍ផ្ទេរទំនិញ", "សាខា")}
    return render_template("print_summary.html", s=data, kind=kind, title=titles[kind][0], who=titles[kind][1])


# ── Sales ────────────────────────────────────────────────────────────────────
@app.get("/sales")
def sales_list():
    items = db.get_all_items()
    return render_template("sales.html", items=items_payload(items), customers=db.get_all_customers(),
                           month_start=month_start(), view=request.args.get("view", ""))


@app.get("/sales/new")
def sale_new():
    return redirect(url_for("sales_list"))


@app.get("/sales/<code>")
def sale_detail(code):
    return redirect(url_for("sales_list", view=code))


def sale_view(r):
    p = pay_row(r["grand_total"], r["paid_amount"])
    return {"code": r["sale_code"], "customer": r["customer_name"], "discount": r["discount_percent"],
            "date": r["sale_date"], **p}


@app.get("/api/sales/list")
def api_sales_list():
    rows = db.get_sales_history(search=(request.args.get("search") or "").strip(),
                                start=parse_date(request.args.get("start")), end=parse_date(request.args.get("end")),
                                paid_filter=request.args.get("paid", "all"))
    return jsonify([sale_view(r) for r in rows])


@app.get("/api/sales/<code>")
def api_sales_get(code):
    r = db.get_sale(code) or abort(404)
    lines = db.get_sale_items(code)
    s1, s2 = split_tel(r["sender_tel"])
    r1, r2 = split_tel(r["receiver_tel"])
    return jsonify(
        code=r["sale_code"], customer=r["customer_name"] or GENERAL_CUSTOMER, date=r["sale_date"], notes=r["notes"] or "",
        sender1=s1, sender2=s2, receiver1=r1, receiver2=r2, discount_percent=r["discount_percent"],
        discount_amount=r["discount_amount"], **pay_row(r["grand_total"], r["paid_amount"]),
        lines=[{"item_id": l["item_id"], "code": l["item_code"], "name": l["name_khmer"] or l["name_english"],
                "qty": l["qty"], "unit_price": l["unit_price"], "sell_price": l["sell_price"],
                "remain": l["remain_qty"] + l["qty"]} for l in lines])


@app.post("/api/sales/save")
def api_sales_save():
    d = body()
    cart = {}
    for x in d.get("items", []):
        qty = safe_float(x.get("qty"))
        if qty <= 0:
            continue
        if x["item_id"] in cart:
            cart[x["item_id"]]["qty"] += qty
        else:
            cart[x["item_id"]] = {"qty": qty, "unit_price": safe_float(x.get("unit_price")),
                                  "sell_price": safe_float(x.get("sell_price"))}
    if not cart:
        return err("សូមបញ្ចូលទំនិញក្នុងកន្ត្រកជាមុនសិន!")
    subtotal = sum(v["qty"] * v["sell_price"] for v in cart.values())
    pct, disc_amt, grand = compute_discounted_total(subtotal, d.get("discount_percent"))
    paid = safe_float(d.get("paid_amount"))
    if paid < 0:
        return err("ប្រាក់ទូទាត់មិនអាចតូចជាងសូន្យបានទេ។")
    editing = (d.get("editing_code") or "").strip()
    customer = (d.get("customer") or "").strip() or GENERAL_CUSTOMER
    try:
        if editing:
            old = db.get_sale(editing) or abort(404)
            code = editing
            sale_date = combine_datetime(d.get("date"), keep_time_from=old["sale_date"])
        else:
            code = db.next_code("INV", "invoice_records", "sale_code")
            sale_date = combine_datetime(d.get("date"))
        db.save_sale(code, sale_date, customer, tel_display(d.get("sender1"), d.get("sender2")),
                     tel_display(d.get("receiver1"), d.get("receiver2")), (d.get("notes") or "").strip(), cart,
                     discount_percent=pct, discount_amount=disc_amt, grand_total=grand, paid_amount=paid,
                     editing=bool(editing))
    except ValueError as e:
        return err(str(e))
    return jsonify(ok=True, code=code, editing=bool(editing), discount_percent=pct, discount_amount=disc_amt,
                   balance=outstanding_balance(grand, paid), grand_total=grand)


@app.post("/api/sales/<code>/payment")
def api_sales_payment(code):
    r = db.get_sale(code) or abort(404)
    new_total = max(safe_float(r["paid_amount"]) + safe_float(body().get("amount")), 0.0)
    db.set_sale_payment(code, new_total)
    return jsonify(ok=True, **pay_row(r["grand_total"], new_total))


@app.post("/api/sales/mark-paid")
def api_sales_mark_paid():
    return jsonify(ok=True, count=db.mark_sales_paid(body().get("codes", [])))


@app.post("/api/sales/<code>/delete")
def api_sales_delete(code):
    db.delete_sale(code)
    return jsonify(ok=True)


@app.post("/api/sales/customer-price")
def api_sales_customer_price():
    d = body()
    if not d.get("customer") or d["customer"] == GENERAL_CUSTOMER:
        return jsonify(ok=True)
    db.save_item_customer_price(d["item_id"], d["customer"], safe_float(d.get("price")))
    return jsonify(ok=True)


@app.get("/api/customers")
def api_customers():
    return jsonify(db.get_all_customers())


@app.post("/api/customers")
def api_customers_add():
    try:
        db.add_customer(body().get("name"))
    except ValueError as e:
        return err(str(e))
    return jsonify(ok=True)


@app.post("/api/customers/rename")
def api_customers_rename():
    d = body()
    try:
        db.rename_customer(d.get("old"), d.get("new"), bool(d.get("update_invoices", True)))
    except ValueError as e:
        return err(str(e))
    return jsonify(ok=True)


@app.get("/sales/<code>/print")
def sale_print(code):
    sale = db.get_sale(code) or abort(404)
    lines = db.get_sale_items(code)
    subtotal = sum(l["qty"] * l["sell_price"] for l in lines)
    return render_template("print_sale.html", sale=sale, lines=lines, subtotal=subtotal,
                           pay=pay_row(sale["grand_total"], sale["paid_amount"]),
                           has_discount=safe_float(sale["discount_percent"]) > MONEY_EPS and subtotal > MONEY_EPS)


# ── Branch transfers ─────────────────────────────────────────────────────────
def transfer_view(r):
    p = pay_row(r["total_cost"], r["paid_amount"])
    return {"code": r["transfer_code"], "branch": r["branch_name"], "driver": r["driver_name"], "date": r["created_at"], **p}


def new_transfer_code():
    return "TRF-" + now_local().strftime("%Y%m%d%H%M%S")


@app.get("/transfers")
def transfers_list():
    items = db.get_all_items()
    return render_template("transfers.html", items=items_payload(items), branches=db.get_branches(),
                           suppliers=db.get_all_suppliers(), month_start=month_start(), view=request.args.get("view", ""))


@app.get("/transfers/new")
def transfer_new():
    return redirect(url_for("transfers_list"))


@app.get("/transfers/<code>")
def transfer_detail(code):
    return redirect(url_for("transfers_list", view=code))


@app.get("/api/transfers/list")
def api_transfers_list():
    rows = db.get_transfers_history(search=(request.args.get("search") or "").strip(),
                                    start=parse_date(request.args.get("start")), end=parse_date(request.args.get("end")),
                                    paid_filter=request.args.get("paid", "all"))
    return jsonify([transfer_view(r) for r in rows])


@app.get("/api/transfers/<code>")
def api_transfers_get(code):
    t = db.get_transfer(code) or abort(404)
    lines = db.get_transfer_items(code)
    return jsonify(
        code=t["transfer_code"], branch=t["branch_name"], driver=t["driver_name"], note=t["note"] or "",
        date=t["created_at"], **pay_row(t["total_cost"], t["paid_amount"]),
        lines=[{"item_id": l["item_id"], "code": l["item_code"], "name": l["name_khmer"] or l["name_english"],
                "qty": l["qty"], "unit_cost": l["unit_cost"], "sell_price": l["sell_price"],
                "supplier": l["supplier_name"], "remark": l["remark"], "remain": l["remain_qty"] + l["qty"]}
               for l in lines])


@app.get("/api/transfers/cost")
def api_transfers_cost():
    """Cost lookup used by the cart: ?item_id=..&supplier=.. (supplier optional)."""
    cost, supplier, src = db.resolve_cost(request.args.get("item_id", ""), (request.args.get("supplier") or "").strip() or None)
    return jsonify(cost=cost, supplier=supplier, source=src)


@app.get("/api/transfers/suppliers")
def api_transfers_suppliers():
    item_id = request.args.get("item_id", "")
    own = db.get_suppliers_for_item(item_id)
    return jsonify(own + [s for s in db.get_all_suppliers() if s not in own])


@app.post("/api/transfers/save")
def api_transfers_save():
    d = body()
    lines = [{"item_id": x["item_id"], "qty": safe_float(x["qty"]), "unit_cost": safe_float(x.get("unit_cost")),
              "sell_price": safe_float(x.get("sell_price", x.get("unit_cost"))), "supplier": x.get("supplier", ""),
              "remark": x.get("remark", "")} for x in d.get("items", []) if safe_float(x.get("qty")) > 0]
    if not lines:
        return err("សូមស្វែងរក/ជ្រើសរើសមុខទំនិញចូលទៅក្នុងកន្ត្រកជាមុនសិន!")
    branch = (d.get("branch") or "").strip()
    if not branch:
        return err("សូមជ្រើសរើសសាខាគោលដៅ!")
    paid = safe_float(d.get("paid_amount"))
    if paid < 0:
        return err("ប្រាក់ទូទាត់មិនអាចតូចជាងសូន្យបានទេ។")
    editing = (d.get("editing_code") or "").strip()
    created = combine_datetime(d.get("date"))
    try:
        code = editing or new_transfer_code()
        total = db.save_transfer(code, branch, (d.get("driver") or "").strip(), (d.get("note") or "").strip(), created,
                                 lines, paid_amount=paid, editing=bool(editing))
    except ValueError as e:
        return err(str(e))
    return jsonify(ok=True, code=code, editing=bool(editing), total=total, balance=outstanding_balance(total, paid))


@app.post("/api/transfers/<code>/payment")
def api_transfers_payment(code):
    t = db.get_transfer(code) or abort(404)
    new_total = max(safe_float(t["paid_amount"]) + safe_float(body().get("amount")), 0.0)
    db.set_transfer_payment(code, new_total)
    return jsonify(ok=True, **pay_row(t["total_cost"], new_total))


@app.post("/api/transfers/mark-paid")
def api_transfers_mark_paid():
    return jsonify(ok=True, count=db.mark_transfers_paid(body().get("codes", [])))


@app.post("/api/transfers/<code>/delete")
def api_transfers_delete(code):
    db.delete_transfer(code)
    return jsonify(ok=True)


@app.get("/transfers/<code>/print")
def transfer_print(code):
    t = db.get_transfer(code) or abort(404)
    lines = db.get_transfer_items(code)
    total = sum(l["qty"] * l["unit_cost"] for l in lines)
    return render_template("print_transfer.html", t=t, lines=lines, total=total, pay=pay_row(total, t["paid_amount"]))


@app.get("/api/branches")
def api_branches():
    return jsonify([{"id": b["id"], "name": b["name"]} for b in db.get_branches()])


@app.post("/api/branches")
def api_branches_add():
    try:
        db.add_branch(body().get("name") or "")
    except ValueError as e:
        return err(str(e))
    return jsonify(ok=True)


@app.post("/api/branches/<branch_id>/delete")
def api_branches_delete(branch_id):
    db.delete_branch(branch_id)
    return jsonify(ok=True)


# ── Labels ───────────────────────────────────────────────────────────────────
@app.get("/labels")
def labels():
    return render_template("labels.html", items=items_payload(db.get_all_items()))


# ── Reports ──────────────────────────────────────────────────────────────────
PAID_FILTERS = {"all": "ទាំងអស់ (All)", "paid": "បានទូទាត់ (Paid)", "unpaid": "មិនទាន់ទូទាត់ (Unpaid)"}
UNKNOWN_NAME = {"customers": "អតិថិជនទូទៅ", "branches": "–", "distributors": "អ្នកផ្គត់ផ្គង់ទូទៅ"}
GROUP_INFO = {
    "distributors": {"entity": "អ្នកផ្គត់ផ្គង់", "title": "របាយការណ៍តាមអ្នកផ្គត់ផ្គង់", "code_label": "វិក្កយបត្រ", "price_label": "ថ្លៃដើម", "icon": "🏭"},
    "customers": {"entity": "អតិថិជន", "title": "របាយការណ៍ទំនាក់ទំនងជាមួយអតិថិជន", "code_label": "វិក្កយបត្រ", "price_label": "តម្លៃលក់", "icon": "👥"},
    "branches": {"entity": "សាខា", "title": "របាយការណ៍ផ្ទេរទំនិញទៅសាខា", "code_label": "ប័ណ្ណផ្ទេរ", "price_label": "ថ្លៃដើម", "icon": "🚚"},
}


def report_params():
    s = parse_date(request.args.get("start"), month_start())
    e = parse_date(request.args.get("end"), today_str())
    if s > e:
        s, e = e, s
    pf = request.args.get("paid", "all")
    return s, e, pf if pf in PAID_FILTERS else "all"


def keep_paid(pf, is_paid):
    return not ((pf == "paid" and not is_paid) or (pf == "unpaid" and is_paid))


def inv_paid_unpaid(cost, paid, is_paid, partial_aware=True):
    cost = safe_float(cost)
    if is_paid:
        return cost, 0.0
    if not partial_aware:
        return 0.0, cost
    p = min(max(safe_float(paid), 0.0), cost)
    return p, cost - p


def group_status(g):
    if g["unpaid"] > 0.001 and g["paid"] > 0.001:
        return "🟡 ទូទាត់ខ្លះ"
    return "❌ មិនទាន់ទូទាត់" if g["unpaid"] > 0.001 else "✅ បានទូទាត់ទាំងអស់"


def build_group(kind, s, e, pf):
    """Distributor / customer / branch report: one group per name, with its invoices."""
    raw = {"distributors": db.report_receives, "customers": db.report_sales, "branches": db.report_transfers}[kind]()
    groups = {}
    for r in raw:
        if not (s <= (r["date"] or "")[:10] <= e):
            continue
        # distributors tab ignores partial payments (desktop behaviour); customers / branches honour them
        paid, unpaid = inv_paid_unpaid(r["total_cost"], r["paid_amount"], r["is_paid"], partial_aware=(kind != "distributors"))
        if not keep_paid(pf, unpaid <= 0.001):
            continue
        name = r.get("distributor_name" if kind == "distributors" else "customer_name" if kind == "customers" else "branch_name") \
            or UNKNOWN_NAME[kind]
        g = groups.setdefault(name, {"invoice_count": 0, "num_items": 0, "total_qty": 0.0, "total": 0.0, "paid": 0.0,
                                     "unpaid": 0.0, "invoices": []})
        g["invoice_count"] += 1
        g["num_items"] += int(r["num_items"] or 0)
        g["total_qty"] += safe_float(r["total_qty"])
        g["total"] += safe_float(r["total_cost"])
        g["paid"] += paid
        g["unpaid"] += unpaid
        g["invoices"].append({"code": r["code"], "date": r["date"], "total": safe_float(r["total_cost"]),
                              "status": "✅ បានទូទាត់" if unpaid <= 0.001 else ("🟡 បង់ខ្លះ" if paid > 0.001 else "❌ មិនទាន់ទូទាត់")})
    return groups


def report_table(tab, s, e, pf):
    """-> dict(cols=[(key,label,align)], rows=[{...}], summary=str, level=...) for one report tab."""
    if tab == "stock":
        rows, tc, ts = [], 0.0, 0.0
        for i in db.get_all_items():
            rem = i["remain_qty"]
            mn, mx, sn, sx = i["min_unit_price"] or 0, i["max_unit_price"] or 0, i["min_sell_price"] or 0, i["max_sell_price"] or 0
            c_tot, s_tot = rem * (i["mid_unit_price"] or 0), rem * (i["mid_sell_price"] or 0)
            tc += c_tot
            ts += s_tot
            sup = ", ".join(dict.fromkeys(x["distributor_name"].strip() for x in i["distributors"] if x["distributor_name"].strip()))
            rows.append({
                "item_code": i["item_code"], "name_khmer": i["name_khmer"] or "–", "name_english": i["name_english"] or "–",
                "floor": i["floor"] or "–", "corner": i["corner"] or "–", "original_qty": fmt_qty(i["original_qty"]),
                "stock_add_qty": fmt_qty(i["stock_add_qty"]), "on_hand_qty": fmt_qty(i["on_hand_qty"]),
                "sell_qty": fmt_qty(i["sell_qty"]), "remain_qty": fmt_qty(rem),
                "cost_range": (fmt_currency(mn) if mn == mx else f"{fmt_currency(mn)} - {fmt_currency(mx)}") if mn else "មិនទាន់កំណត់",
                "distributors": sup or "–",
                "sell_range": (fmt_currency(sn) if sn == sx else f"{fmt_currency(sn)} - {fmt_currency(sx)}") if sn else "មិនទាន់កំណត់",
                "total_unit_price": fmt_currency(c_tot), "total_sell_price": fmt_currency(ts and s_tot)})
        cols = [("item_code", "កូដទំនិញ", "w"), ("name_khmer", "ឈ្មោះខ្មែរ", "w"), ("name_english", "ឈ្មោះអង់គ្លេស", "w"),
                ("floor", "ជាន់", "c"), ("corner", "ជួរ", "c"), ("original_qty", "បរិមាណដើមដំបូង", "e"),
                ("stock_add_qty", "បរិមាណបញ្ចូលថ្មី", "e"), ("on_hand_qty", "ស្តុកសរុប", "e"), ("sell_qty", "លក់ចេញ", "e"),
                ("remain_qty", "នៅសល់", "e"), ("cost_range", "ថ្លៃដើម (ទាប-ខ្ពស់)", "e"), ("distributors", "អ្នកផ្គត់ផ្គង់", "w"),
                ("sell_range", "តម្លៃលក់ (ទាប-ខ្ពស់)", "e"), ("total_unit_price", "តម្លៃដើមសរុប (Avg)", "e"),
                ("total_sell_price", "តម្លៃលក់សរុប (Avg)", "e")]
        return dict(cols=cols, rows=rows, level="accent", summary=(
            f"📊 បញ្ជីសារពើភ័ណ្ឌស្តុកសរុបរួម ({s} ដល់ {e})\n• មុខទំនិញសរុប: {len(rows)} មុខ    "
            f"• ថ្លៃដើមវាយតម្លៃស្តុកសរុប (Avg): {fmt_currency(tc)}    • តម្លៃទីផ្សារប៉ាន់ស្មានសរុប (Avg): {fmt_currency(ts)}"))

    if tab == "receives":
        rows, tot, unp = [], 0.0, 0.0
        for r in db.report_receives():
            if not (s <= (r["date"] or "")[:10] <= e) or not keep_paid(pf, r["is_paid"]):
                continue
            c = safe_float(r["total_cost"])
            tot += c
            unp += 0 if r["is_paid"] else c
            rows.append({"receive_code": r["code"], "distributor_name": r["distributor_name"] or UNKNOWN_NAME["distributors"],
                         "num_items": str(r["num_items"]), "total_qty": fmt_qty(r["total_qty"]), "receive_date": r["date"],
                         "paid_status": "✅ បានទូទាត់" if r["is_paid"] else "❌ មិនទាន់ទូទាត់", "notes": r["notes"] or "–",
                         "_code": r["code"]})
        cols = [("receive_code", "លេខកូដនាំចូល", "w"), ("distributor_name", "អ្នកផ្គត់ផ្គង់/ចែកចាយ", "w"),
                ("num_items", "ចំនួនមុខទំនិញ", "c"), ("total_qty", "បរិមាណសរុបរួម", "c"), ("receive_date", "កាលបរិច្ឆេទនាំចូល", "c"),
                ("paid_status", "ស្ថានភាពទូទាត់", "c"), ("notes", "កំណត់សម្គាល់", "w")]
        return dict(cols=cols, rows=rows, level="danger" if unp > 0.001 else "ok", print_kind="receive", summary=(
            f"📥 របាយការណ៍នាំចូលទំនិញពីអ្នកផ្គត់ផ្គង់ ({s} ដល់ {e})\n• វិក្កយបត្រសរុប: {len(rows)} ច្បាប់    "
            f"• ទឹកប្រាក់ចំណាយទិញចូលសរុប: {fmt_currency(tot)}    • មិនទាន់ទូទាត់ទៅអ្នកផ្គត់ផ្គង់: {fmt_currency(unp)}"))

    if tab == "sales":
        rows, tot, unp, disc = [], 0.0, 0.0, 0.0
        for r in db.report_sales():
            if not (s <= (r["date"] or "")[:10] <= e) or not keep_paid(pf, r["is_paid"]):
                continue
            c = safe_float(r["total_cost"])
            tot += c
            disc += safe_float(r["discount_amount"])
            unp += 0 if r["is_paid"] else c
            note = r["notes"] or ""
            if "|" in note and len(note.split("|")) >= 4:
                note = note.split("|")[3].strip() or "–"
            rows.append({"sale_code": r["code"], "customer_name": r["customer_name"] or UNKNOWN_NAME["customers"],
                         "num_items": str(r["num_items"]), "total_qty": fmt_qty(r["total_qty"]),
                         "discount": fmt_pct(r["discount_percent"]) if safe_float(r["discount_percent"]) > 0.001 else "–",
                         "total_cost": fmt_currency(c), "paid_status": "✅ បានទូទាត់" if r["is_paid"] else "❌ មិនទាន់ទូទាត់",
                         "sale_date": r["date"], "notes": note if note.strip() else "–", "_code": r["code"]})
        cols = [("sale_code", "លេខវិក្កយបត្រ", "w"), ("customer_name", "ឈ្មោះអតិថិជន", "w"), ("num_items", "ចំនួនមុខទំនិញ", "c"),
                ("total_qty", "បរិមាណសរុបរួម", "c"), ("discount", "បញ្ចុះតម្លៃ", "c"), ("total_cost", "ថ្លៃលក់សរុប", "e"),
                ("paid_status", "ស្ថានភាពទូទាត់", "c"), ("sale_date", "កាលបរិច្ឆេទលក់", "c"), ("notes", "កំណត់សម្គាល់", "w")]
        return dict(cols=cols, rows=rows, level="danger" if unp > 0.001 else "ok", print_kind="sale", summary=(
            f"📤 របាយការណ៍លក់ ផ្អែកលើវិក្កយបត្រសរុប ({s} ដល់ {e})\n• វិក្កយបត្រសរុប: {len(rows)} ច្បាប់    "
            f"• ទឹកប្រាក់លក់សរុប (ក្រោយបញ្ចុះតម្លៃ): {fmt_currency(tot)}    • សរុបបញ្ចុះតម្លៃ: {fmt_currency(disc)}    "
            f"• មិនទាន់ទូទាត់ពីអតិថិជន: {fmt_currency(unp)}"))

    if tab == "transfers":
        rows, tot, unp = [], 0.0, 0.0
        for r in db.report_transfers():
            if not (s <= (r["date"] or "")[:10] <= e) or not keep_paid(pf, r["is_paid"]):
                continue
            c = safe_float(r["total_cost"])
            tot += c
            if not r["is_paid"]:
                unp += max(c - safe_float(r["paid_amount"]), 0.0)
            rows.append({"transfer_code": r["code"], "branch_name": r["branch_name"] or "–", "driver_name": r["driver_name"] or "–",
                         "item_count": str(r["num_items"]), "total_qty": fmt_qty(r["total_qty"]), "total_cost": fmt_currency(c),
                         "created_at": r["date"], "paid_status": "✅ បានទូទាត់" if r["is_paid"] else "❌ មិនទាន់ទូទាត់",
                         "note": r["note"] or "–", "_code": r["code"]})
        cols = [("transfer_code", "លេខប័ណ្ណផ្ទេរ", "w"), ("branch_name", "សាខាគោលដៅ", "w"), ("driver_name", "អ្នកដឹកជញ្ជូន", "w"),
                ("item_count", "ចំនួនមុខទំនិញ", "c"), ("total_qty", "បរិមាណសរុបរួម", "c"), ("total_cost", "ថ្លៃដើមសរុប", "e"),
                ("created_at", "កាលបរិច្ឆេទផ្ទេរ", "c"), ("paid_status", "ស្ថានភាពទូទាត់", "c"), ("note", "កំណត់សម្គាល់", "w")]
        return dict(cols=cols, rows=rows, level="danger" if unp > 0.001 else "ok", print_kind="transfer", summary=(
            f"🚚 របាយការណ៍ផ្ទេរទំនិញរវាងសាខា ({s} ដល់ {e})\n• ប័ណ្ណផ្ទេរសរុប: {len(rows)} ច្បាប់    "
            f"• ថ្លៃដើមផ្ទេរសរុប: {fmt_currency(tot)}    • មិនទាន់ទូទាត់: {fmt_currency(unp)}"))

    if tab in GROUP_INFO:
        info = GROUP_INFO[tab]
        groups = build_group(tab, s, e, pf)
        rows, unp_all = [], 0.0
        for name, g in sorted(groups.items(), key=lambda kv: kv[0]):
            unp_all += g["unpaid"]
            rows.append({"name": name, "invoice_count": str(g["invoice_count"]), "num_items": str(g["num_items"]),
                         "total_qty": fmt_qty(g["total_qty"]), "total_cost": fmt_currency(g["total"]),
                         "paid_amount": fmt_currency(g["paid"]), "unpaid_amount": fmt_currency(g["unpaid"]),
                         "status": group_status(g), "_name": name})
        cols = [("name", info["entity"], "w"), ("invoice_count", "ចំនួនវិក្កយបត្រ", "c"), ("num_items", "ចំនួនមុខទំនិញសរុប", "c"),
                ("total_qty", "បរិមាណសរុប", "c"), ("total_cost", "ទឹកប្រាក់សរុប", "e"), ("paid_amount", "បានទូទាត់", "e"),
                ("unpaid_amount", "មិនទាន់ទូទាត់", "e"), ("status", "ស្ថានភាព", "c")]
        return dict(cols=cols, rows=rows, level="danger" if unp_all > 0.001 else "ok", group=tab, summary=(
            f"{info['icon']} {info['title']} ({s} ដល់ {e})\n• {info['entity']}ដែលមានប្រតិបត្តិការ: {len(rows)}    "
            f"• ទឹកប្រាក់មិនទាន់ទូទាត់សរុប: {fmt_currency(unp_all)}    • ចុចទ្វេដងលើជួរ ដើម្បីមើលមុខទំនិញលម្អិត"))

    if tab == "profit":
        lines = db.sales_report(s, e)
        tot = db.sales_report_totals(s, e)
        gross = sum(l["qty"] * l["sell_price"] for l in lines)
        cost = sum(l["qty"] * l["unit_price"] for l in lines)
        revenue = safe_float(tot["revenue"])
        profit, margin = calc_profit(revenue, cost)
        top = {}
        for l in lines:
            t = top.setdefault(l["item_id"], {"code": l["item_code"], "name": l["name"], "qty": 0, "revenue": 0, "cost": 0})
            t["qty"] += l["qty"]
            t["revenue"] += l["qty"] * l["sell_price"]
            t["cost"] += l["qty"] * l["unit_price"]
        rows = [{"item_code": t["code"], "name": t["name"], "qty": fmt_qty(t["qty"]), "revenue": fmt_currency(t["revenue"]),
                 "profit": fmt_currency(t["revenue"] - t["cost"])}
                for t in sorted(top.values(), key=lambda r: r["revenue"], reverse=True)[:50]]
        cols = [("item_code", "កូដទំនិញ", "w"), ("name", "ឈ្មោះ", "w"), ("qty", "បរិមាណលក់", "e"), ("revenue", "ចំណូល", "e"), ("profit", "ប្រាក់ចំណេញ", "e")]
        return dict(cols=cols, rows=rows, level="ok" if profit >= 0 else "danger", summary=(
            f"💹 ប្រាក់ចំណេញ ({s} ដល់ {e})\n• ចំណូលលក់ (ក្រោយបញ្ចុះតម្លៃ): {fmt_currency(revenue)}    • ថ្លៃដើមទំនិញលក់: {fmt_currency(cost)}    "
            f"• ប្រាក់ចំណេញ: {fmt_currency(profit)} ({margin}%)    • បញ្ចុះតម្លៃសរុប: {fmt_currency(tot['discount'])}"))
    abort(404)


@app.get("/reports")
def reports():
    return render_template("reports.html", start=month_start(), end=today_str(), paid_filters=PAID_FILTERS)


@app.get("/api/reports/<tab>")
def api_reports(tab):
    s, e, pf = report_params()
    r = report_table(tab, s, e, pf)
    r["cols"] = [list(c) for c in r["cols"]]
    return jsonify(r)


@app.get("/api/reports/group/<kind>")
def api_reports_group(kind):
    if kind not in GROUP_INFO:
        abort(404)
    s, e, pf = report_params()
    name = request.args.get("name", "")
    g = build_group(kind, s, e, pf).get(name) or abort(404)
    inv = []
    for iv in g["invoices"]:
        lines = db.report_invoice_lines(kind, iv["code"])
        inv.append({**iv, "lines": [{"item_code": l["item_code"], "name": l["name_khmer"] or l["item_code"], "qty": l["qty"],
                                      "price": l["price"], "total": l["qty"] * l["price"]} for l in lines]})
    return jsonify(name=name, info=GROUP_INFO[kind], start=s, end=e, invoices=inv,
                   **{k: v for k, v in g.items() if k != "invoices"})


@app.get("/reports/group/print")
def reports_group_print():
    kind = request.args.get("kind", "")
    if kind not in GROUP_INFO:
        abort(404)
    data = api_reports_group(kind).get_json()
    return render_template("print_group.html", g=data, kind=kind)


@app.get("/reports/export")
def reports_export():
    """Excel export of the active report tab (.xlsx)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    tab = request.args.get("tab", "stock")
    s, e, pf = report_params()
    wb = Workbook()
    ws = wb.active
    title = tab
    if tab == "receives":
        title = "receives_detail"
        head = ["លេខកូដនាំចូល", "កាលបរិច្ឆេទនាំចូល", "អ្នកផ្គត់ផ្គង់/ចែកចាយ", "ស្ថានភាពទូទាត់", "កូដទំនិញ", "ឈ្មោះទំនិញ (ខ្មែរ)",
                "បរិមាណនាំចូល", "ថ្លៃដើម/ឯកតា", "សរុបទឹកប្រាក់", "កំណត់សម្គាល់"]
        data = []
        for r in db.report_receives():
            if not (s <= (r["date"] or "")[:10] <= e) or not keep_paid(pf, r["is_paid"]):
                continue
            for l in db.report_invoice_lines("distributors", r["code"]):
                data.append([r["code"], r["date"], r["distributor_name"] or UNKNOWN_NAME["distributors"],
                             "បានទូទាត់" if r["is_paid"] else "មិនទាន់ទូទាត់", l["item_code"], l["name_khmer"] or "–",
                             l["qty"], l["price"], l["qty"] * l["price"], r["notes"] or "–"])
    elif tab == "sales":
        title = "sales_detail"
        head = ["លេខវិក្កយបត្រ", "កាលបរិច្ឆេទលក់", "ឈ្មោះអតិថិជន", "ស្ថានភាពទូទាត់", "បញ្ចុះតម្លៃ (%)", "ទឹកប្រាក់បញ្ចុះតម្លៃ",
                "លេខអ្នកផ្ញើ", "លេខអ្នកទទួល", "កូដទំនិញ", "ឈ្មោះទំនិញ (ខ្មែរ)", "ឈ្មោះទំនិញ (អង់គ្លេស)", "បរិមាណ", "តម្លៃលក់/ឯកតា",
                "សរុបទឹកប្រាក់", "កំណត់សម្គាល់"]
        data = []
        for r in db.report_sales():
            if not (s <= (r["date"] or "")[:10] <= e) or not keep_paid(pf, r["is_paid"]):
                continue
            for l in db.get_sale_items(r["code"]):
                data.append([r["code"], r["date"], r["customer_name"] or UNKNOWN_NAME["customers"],
                             "បានទូទាត់" if r["is_paid"] else "មិនទាន់ទូទាត់", r["discount_percent"], r["discount_amount"],
                             r["sender_tel"], r["receiver_tel"], l["item_code"], l["name_khmer"] or "–", l["name_english"] or "–",
                             l["qty"], l["sell_price"], l["qty"] * l["sell_price"], r["notes"] or "–"])
    elif tab == "transfers":
        title = "transfers_detail"
        head = ["លេខប័ណ្ណផ្ទេរ", "កាលបរិច្ឆេទផ្ទេរ", "សាខាគោលដៅ", "អ្នកដឹកជញ្ជូន", "ស្ថានភាពទូទាត់", "កូដទំនិញ", "ឈ្មោះទំនិញ (ខ្មែរ)",
                "បរិមាណ", "ថ្លៃដើម/ឯកតា", "តម្លៃលក់/ឯកតា", "អ្នកផ្គត់ផ្គង់", "សម្គាល់", "សរុបថ្លៃដើម", "កំណត់សម្គាល់ប័ណ្ណ"]
        data = []
        for r in db.report_transfers():
            if not (s <= (r["date"] or "")[:10] <= e) or not keep_paid(pf, r["is_paid"]):
                continue
            for l in db.get_transfer_items(r["code"]):
                data.append([r["code"], r["date"], r["branch_name"], r["driver_name"], "បានទូទាត់" if r["is_paid"] else "មិនទាន់ទូទាត់",
                             l["item_code"], l["name_khmer"] or "–", l["qty"], l["unit_cost"], l["sell_price"], l["supplier_name"],
                             l["remark"], l["qty"] * l["unit_cost"], r["note"] or "–"])
    else:
        t = report_table(tab, s, e, pf)
        head = [c[1] for c in t["cols"]]
        data = [[r.get(c[0], "") for c in t["cols"]] for r in t["rows"]]
    ws.title = title[:31]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2563EB")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in data:
        ws.append(row)
    for i, col in enumerate(ws.columns, 1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 45)
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"report_{tab}_{s}_{e}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ── Settings ─────────────────────────────────────────────────────────────────
@app.route("/settings", methods=["GET", "POST"])
def settings():
    if request.method == "POST":
        try:
            db.add_branch(f("name"))
            flash("បានបន្ថែមសាខារួចរាល់។", "ok")
        except ValueError as e:
            flash(str(e), "error")
        return redirect(url_for("settings"))
    return render_template("settings.html", branches=db.get_branches(), author=APP_AUTHOR, year=now_local().year,
                           drive_ok=bk.drive_configured(), cron_ready=bool(os.environ.get("CRON_SECRET")),
                           keep=os.environ.get("BACKUP_KEEP", "30"), runs=bk.recent_runs(8),
                           folder=os.environ.get("GDRIVE_FOLDER_NAME", "") or "LKN Stock Backups")


@app.post("/settings/branch/<branch_id>/delete")
def branch_delete(branch_id):
    db.delete_branch(branch_id)
    flash("បានលុបសាខារួចរាល់។", "ok")
    return redirect(url_for("settings"))


@app.post("/settings/backup")
def settings_backup():
    """Download the whole database: one .db file (default, SQLite format) or a .json file."""
    if not check_password(request.form.get("password", "")):
        flash("ពាក្យសម្ងាត់មិនត្រឹមត្រូវ — មិនបានបង្កើតទិន្នន័យបម្រុង។", "error")
        return redirect(url_for("settings"))
    if request.form.get("format") == "json":
        payload = {"app": "lknstock", "version": APP_VERSION, "created": now_local().strftime("%Y-%m-%d %H:%M:%S"),
                   **db.export_all()}
        raw = json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8")
        name = "stock_backup_" + now_local().strftime("%Y-%m-%d_%H-%M-%S") + ".json"
        return send_file(io.BytesIO(raw), as_attachment=True, download_name=name, mimetype="application/json")
    try:
        name, raw, _ = bk.make_backup_file(APP_VERSION)
    except Exception as e:
        flash(f"មិនអាចបង្កើតឯកសារបម្រុងបានទេ៖ {e}", "error")
        return redirect(url_for("settings"))
    return send_file(io.BytesIO(raw), as_attachment=True, download_name=name, mimetype="application/x-sqlite3")


@app.post("/settings/backup-drive")
def settings_backup_drive():
    """'Back up to Google Drive now' button (needs the normal login)."""
    return jsonify(bk.run_backup("manual", APP_VERSION))


@app.route("/cron/backup", methods=["GET", "POST"])
def cron_backup():
    """Called by the scheduler (cron-job.org etc.) every evening. Protected by CRON_SECRET, not by the login page.
    Send the secret as 'Authorization: Bearer <secret>' (preferred) or ?key=<secret>.
    Returns 200 on success, 500 if the backup failed (so the scheduler can alert you)."""
    # tolerate the usual paste mistakes: quotes / spaces around the value, "Bearer" typed twice, secret in header OR in ?key=
    clean = lambda v: (v or "").strip().strip("\"'").strip()
    secret = clean(os.environ.get("CRON_SECRET", ""))
    auth = clean(request.headers.get("Authorization", ""))
    candidates = [clean(auth[7:]) if auth.lower().startswith("bearer ") else auth,
                  clean(request.headers.get("X-Cron-Secret", "")), clean(request.args.get("key", ""))]
    candidates = [c for c in candidates if c]
    if not secret or not any(hmac.compare_digest(c.encode(), secret.encode()) for c in candidates):
        # say WHICH part is wrong (never the secret itself) so the cron job can be fixed quickly
        why = ("CRON_SECRET is not set on the server (add it in Render > Environment, then redeploy)" if not secret
               else "no secret was sent (add the header 'Authorization: Bearer <CRON_SECRET>' or '?key=<CRON_SECRET>' to the URL)" if not candidates
               else "the secret that was sent does not match CRON_SECRET (check for typos / an old value; redeploy after changing it)")
        return jsonify(ok=False, error="forbidden", hint=why), 403
    result = bk.run_backup("cron", APP_VERSION)
    return jsonify(result), (200 if result.get("ok") else 500)


@app.post("/settings/restore")
def settings_restore():
    if not check_password(request.form.get("password", "")):
        flash("ពាក្យសម្ងាត់មិនត្រឹមត្រូវ — មិនបានស្តារទិន្នន័យ។", "error")
        return redirect(url_for("settings"))
    file = request.files.get("backup")
    if not file or not file.filename:
        flash("សូមជ្រើសរើសឯកសារទិន្នន័យបម្រុង (.db ឬ .json)។", "error")
        return redirect(url_for("settings"))
    try:
        raw = file.read()
        data = bk.read_sqlite_backup(raw) if raw.startswith(bk.SQLITE_MAGIC) else json.loads(raw.decode("utf-8"))
        counts = db.restore_all(data)
        flash(f"ទិន្នន័យត្រូវបានស្តារឡើងវិញរួចរាល់ ({counts['items']} ទំនិញ, {counts['invoice_records']} វិក្កយបត្រលក់, "
              f"{counts['purchase_invoices']} វិក្កយបត្រទទួលស្តុក, {counts['transfers']} ប័ណ្ណផ្ទេរ)។", "ok")
    except (json.JSONDecodeError, UnicodeDecodeError):
        flash("ឯកសារនេះមិនមែនជាឯកសារបម្រុង (.db ឬ .json) របស់កម្មវិធីនេះទេ។", "error")
    except Exception as e:
        flash(f"មិនអាចស្តារទិន្នន័យបានទេ ដោយសារតែ៖ {e}", "error")
    return redirect(url_for("settings"))


@app.post("/settings/reset")
def settings_reset():
    d = body()
    if not check_password(d.get("password", "")):
        return err("ពាក្យសម្ងាត់មិនត្រឹមត្រូវ។", 403)
    if (d.get("word") or "").strip() != "DELETE":
        return err("សូមវាយពាក្យ DELETE ឱ្យត្រឹមត្រូវ។")
    cleared = db.reset_all()
    return jsonify(ok=True, cleared=len(cleared))


@app.errorhandler(413)
def too_large(_):
    flash("ឯកសារធំពេក (File too large).", "error")
    return redirect(request.referrer or url_for("items"))


if __name__ == "__main__":
    app.run(debug=True, port=int(os.environ.get("PORT", 5000)))
